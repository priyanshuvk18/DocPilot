"""
Document parser supporting PDF, DOCX, TXT, and Markdown.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

import structlog

logger = structlog.get_logger(__name__)


@dataclass
class ParsedPage:
    page_number: int
    text: str


@dataclass
class ParsedDocument:
    filename: str
    pages: list[ParsedPage]
    total_pages: int


def parse_pdf(content: bytes, filename: str) -> ParsedDocument:
    """Parse PDF bytes using PyMuPDF (fitz) with fallback to pypdf."""
    pages: list[ParsedPage] = []
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=content, filetype="pdf")
        for i, page in enumerate(doc):
            text = page.get_text()
            pages.append(ParsedPage(page_number=i + 1, text=text))
        return ParsedDocument(filename=filename, pages=pages, total_pages=len(pages))
    except Exception as e:
        logger.warning("PyMuPDF failed, falling back to pypdf", error=str(e))

    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content))
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append(ParsedPage(page_number=i + 1, text=text))
        return ParsedDocument(filename=filename, pages=pages, total_pages=len(pages))
    except Exception as e:
        logger.error("Failed to parse PDF", error=str(e))
        return ParsedDocument(filename=filename, pages=[ParsedPage(1, "")], total_pages=1)


def parse_docx(content: bytes, filename: str) -> ParsedDocument:
    """Parse DOCX bytes using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        full_text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return ParsedDocument(
            filename=filename,
            pages=[ParsedPage(page_number=1, text=full_text)],
            total_pages=1,
        )
    except Exception as e:
        logger.error("Failed to parse DOCX", error=str(e))
        return ParsedDocument(filename=filename, pages=[ParsedPage(1, "")], total_pages=1)


def parse_text(content: bytes, filename: str) -> ParsedDocument:
    """Parse plain text or Markdown with automatic encoding detection."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = content.decode(encoding)
            return ParsedDocument(
                filename=filename,
                pages=[ParsedPage(page_number=1, text=text)],
                total_pages=1,
            )
        except UnicodeDecodeError:
            continue

    text = content.decode("utf-8", errors="replace")
    return ParsedDocument(
        filename=filename,
        pages=[ParsedPage(page_number=1, text=text)],
        total_pages=1,
    )


def parse_document(content: bytes, filename: str) -> ParsedDocument:
    """Route document to appropriate parser by file extension."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(content, filename)
    elif ext in (".docx", ".doc"):
        return parse_docx(content, filename)
    else:
        return parse_text(content, filename)
